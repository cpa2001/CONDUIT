#!/usr/bin/env python3

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import Callable, cast

from export_sweep_metrics_csv import (
    consolidate_export_records,
    extract_ratio_from_label,
    format_value,
    format_ratio_percent,
    get_benchmark_name,
    get_export_records,
    get_metric_entry,
    get_record_image_priori_mode,
    get_record_input_max_length,
    get_record_test_file,
    load_json,
    normalize_row_label,
    parse_row_label,
    should_include_record,
)


DEFAULT_BASELINE_TAG = "all_configs_8k_sweep"
DEFAULT_EXPERIMENT_TAG = "experiment_8k"
DEFAULT_IMAGE_PRIORI_MODE = "chat_template"
DEFAULT_PREFILL_MODE = "image_segment"
DEFAULT_CASES = "vanilla"
DEFAULT_INFOBLEND_TAG = "infoblend_experiment_8k"
DEFAULT_CACHEBLEND_TAG = "cacheblend_experiment_8k"
DEFAULT_KVSHARE_TAG = "kvshare_experiment_8k"
DEFAULT_NONE_TAG = "none-experiment_8k"
DEFAULT_NONE_IMAGE_PRIORI_MODE = "none"
DEFAULT_METHOD_TOKEN_ALIASES = "imgbias100=reweight,scoresplit4of4=split4of4"
DEFAULT_BENCHMARK_TAIL = "vh_single,vh_multi"
DEFAULT_METHOD_VARIANT_ORDER = (
    "base,"
    "vnorm,"
    "reweight,"
    "vnorm+reweight,"
    "split4of4,"
    "split4of4+vnorm,"
    "split4of4+reweight,"
    "split4of4+vnorm+reweight"
)

CASE_SUFFIX_ALIASES = {
    "vanilla": "",
    "split4": "scoresplit4of4",
    "last1": "scorelast1",
}

SPECIAL_METHOD_ORDER = {
    "full": 0,
    "image_segment": 1,
    "none": 2,
}

STRATEGY_DISPLAY_ORDER = {
    "infoblend": 0,
    "cacheblend": 1,
    "kvshare": 2,
    "kv_score": 3,
}

PREFERRED_MODEL_ORDER = (
    "qwen2.5-vl-3b-instruct",
    "qwen2.5-vl-7b-instruct",
    "internvl3-9b",
)

NON_ALNUM_PATTERN = re.compile(r"[^a-z0-9]+")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Export the combined MMLongBench KV score total experiment table, "
            "including optional full/image_segment special baselines, the shared "
            "none baseline, InfoBlend, CacheBlend, KVShare, and KV score case "
            "summaries. When multiple models are selected, the CSV/XLSX is "
            "written as one combined table with model-separated blocks."
        )
    )
    parser.add_argument(
        "--benchmark-root",
        type=Path,
        default=Path(__file__).resolve().parent.parent,
        help="MMLongBench root directory. Defaults to the parent of this scripts directory.",
    )
    parser.add_argument(
        "--model-output-dir",
        action="append",
        default=[],
        help=(
            "Model output directory name under output/, for example "
            "qwen2.5-vl-3b-instruct. Repeatable; comma-separated values are also "
            "accepted. When omitted, all matching model output directories are "
            "auto-discovered."
        ),
    )
    parser.add_argument(
        "--baseline-tag",
        default=DEFAULT_BASELINE_TAG,
        help=(
            "Legacy fallback summary tag used for the special full/image_segment "
            "baseline rows. Pass an empty string to disable both when no per-baseline "
            "tag is provided."
        ),
    )
    parser.add_argument(
        "--full-baseline-tag",
        default=None,
        help=(
            "Summary tag used for the special full baseline row. Defaults to "
            "--baseline-tag. Pass an empty string to skip the full baseline row."
        ),
    )
    parser.add_argument(
        "--segment-baseline-tag",
        default=None,
        help=(
            "Summary tag used for the special image_segment baseline row. Defaults to "
            "--baseline-tag. Pass an empty string to skip the image_segment baseline row."
        ),
    )
    parser.add_argument(
        "--experiment-tag",
        default=DEFAULT_EXPERIMENT_TAG,
        help="Base tag of the KV score experiment.",
    )
    parser.add_argument(
        "--image-priori-mode",
        default=DEFAULT_IMAGE_PRIORI_MODE,
        help="Image priori mode used by the KV score, InfoBlend, CacheBlend, and KVShare summaries.",
    )
    parser.add_argument(
        "--prefill-mode",
        default=DEFAULT_PREFILL_MODE,
        help="Prefill mode used by the KV score, InfoBlend, CacheBlend, and KVShare summaries.",
    )
    parser.add_argument(
        "--cases",
        default=DEFAULT_CASES,
        help=(
            "Comma-separated case names or suffixes. Built-in aliases: split4 -> "
            "scoresplit4of4, last1 -> scorelast1. Ignored when --case-summary is used."
        ),
    )
    parser.add_argument(
        "--case-summary",
        action="append",
        type=Path,
        default=[],
        help=(
            "Explicit path to a case summary JSON. Repeatable. When provided, these "
            "paths are used directly instead of resolving --cases. This option is "
            "only supported when exporting exactly one model."
        ),
    )
    parser.add_argument(
        "--shared-summary-root",
        type=Path,
        help=(
            "Shared summary root used for cross-model summaries such as none-experiment_8k. "
            "Defaults to <benchmark-root>/output/_summaries."
        ),
    )
    parser.add_argument(
        "--infoblend-tag",
        default=DEFAULT_INFOBLEND_TAG,
        help=(
            "Summary tag used for the InfoBlend rows. Pass an empty string to disable "
            "Infoblend row export."
        ),
    )
    parser.add_argument(
        "--cacheblend-tag",
        default=DEFAULT_CACHEBLEND_TAG,
        help=(
            "Summary tag used for the CacheBlend rows. Pass an empty string to disable "
            "CacheBlend row export."
        ),
    )
    parser.add_argument(
        "--kvshare-tag",
        default=DEFAULT_KVSHARE_TAG,
        help=(
            "Summary tag used for the KVShare rows. Pass an empty string to disable "
            "KVShare row export."
        ),
    )
    parser.add_argument(
        "--none-tag",
        default=DEFAULT_NONE_TAG,
        help=(
            "Shared summary tag used for the no-priori vanilla rows. Pass an empty string "
            "to disable none row export."
        ),
    )
    parser.add_argument(
        "--none-image-priori-mode",
        default=DEFAULT_NONE_IMAGE_PRIORI_MODE,
        help="Image priori mode used by the shared none summary.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help=(
            "Output CSV path. Defaults to output/_summaries/<experiment-tag>/benchmark_metric_table.csv "
            "for auto-discovered multi-model export, or to "
            "output/<model-output-dir>/_summaries/<experiment-tag>/benchmark_metric_table.csv "
            "when a single model is explicitly requested."
        ),
    )
    parser.add_argument(
        "--xlsx-output",
        type=Path,
        help="Output XLSX path. Defaults to the CSV output path with a .xlsx suffix.",
    )
    parser.add_argument(
        "--method-token-aliases",
        default=DEFAULT_METHOD_TOKEN_ALIASES,
        help=(
            "Comma-separated token aliases applied to method suffixes in the output table. "
            "Use old=new entries, for example imgbias100=reweight. "
            "Pass an empty string to disable aliases."
        ),
    )
    parser.add_argument(
        "--benchmark-tail",
        default=DEFAULT_BENCHMARK_TAIL,
        help=(
            "Comma-separated benchmark names to move to the end of the table columns, "
            "preserving their relative order. Pass an empty string to keep discovery order."
        ),
    )
    parser.add_argument(
        "--method-variant-order",
        default=DEFAULT_METHOD_VARIANT_ORDER,
        help=(
            "Comma-separated variant combinations used to sort KV score method rows "
            "after aliasing. Use 'base' for rows without extra suffix tokens and '+' "
            "to combine tokens, for example vnorm+reweight."
        ),
    )
    return parser.parse_args()


def split_csv(value: str) -> list[str]:
    return [item.strip() for item in value.split(",") if item.strip()]


def parse_assignment_csv(value: str) -> dict[str, str]:
    assignments: dict[str, str] = {}
    for item in split_csv(value):
        source, separator, target = item.partition("=")
        source = source.strip()
        target = target.strip()
        if not separator or not source:
            raise ValueError(
                "Expected alias entries in old=new form, got "
                f"{item!r} from --method-token-aliases"
            )
        assignments[source] = target
    return assignments


def parse_variant_order_csv(value: str) -> tuple[tuple[str, ...], ...]:
    variant_order: list[tuple[str, ...]] = []
    seen: set[frozenset[str]] = set()
    for item in split_csv(value):
        if item == "base":
            tokens: tuple[str, ...] = ()
        else:
            tokens = tuple(token.strip() for token in item.split("+") if token.strip())
            if not tokens:
                raise ValueError(
                    "Expected non-empty variant tokens in --method-variant-order, "
                    f"got {item!r}"
                )

        combo_key = frozenset(tokens)
        if combo_key in seen:
            raise ValueError(
                "Duplicate variant combination in --method-variant-order: "
                f"{item!r}"
            )
        seen.add(combo_key)
        variant_order.append(tokens)

    return tuple(variant_order)


def resolve_configured_summary_tag(configured_tag: str | None, fallback_tag: str) -> str:
    if configured_tag is None:
        return fallback_tag.strip()
    return configured_tag.strip()


def split_method_annotation(row_label: str) -> tuple[str, str]:
    if row_label.endswith("]"):
        annotation_start = row_label.rfind(" [")
        if annotation_start != -1:
            return row_label[:annotation_start], row_label[annotation_start:]
    return row_label, ""


def split_variant_tokens(variant_suffix: str) -> list[str]:
    return [token for token in variant_suffix.split("-") if token]


def apply_method_token_aliases(
    tokens: list[str],
    token_aliases: dict[str, str],
) -> list[str]:
    aliased_tokens: list[str] = []
    for token in tokens:
        mapped_token = token_aliases.get(token, token)
        if mapped_token:
            aliased_tokens.append(mapped_token)
    return aliased_tokens


def build_variant_order_lookup(
    variant_order: tuple[tuple[str, ...], ...],
) -> dict[frozenset[str], tuple[int, tuple[str, ...]]]:
    lookup: dict[frozenset[str], tuple[int, tuple[str, ...]]] = {}
    for index, combo in enumerate(variant_order):
        lookup[frozenset(combo)] = (index, combo)
    return lookup


def build_variant_token_priority(
    variant_order: tuple[tuple[str, ...], ...],
) -> dict[str, int]:
    priority: dict[str, int] = {}
    for combo in variant_order:
        for token in combo:
            priority.setdefault(token, len(priority))
    return priority


def resolve_variant_order_entry(
    tokens: list[str],
    variant_order_lookup: dict[frozenset[str], tuple[int, tuple[str, ...]]],
) -> tuple[int, tuple[str, ...]] | None:
    return variant_order_lookup.get(frozenset(tokens))


def reorder_benchmarks(
    benchmark_order: list[str],
    benchmark_tail: tuple[str, ...],
) -> list[str]:
    if not benchmark_tail:
        return list(benchmark_order)

    benchmark_tail_set = set(benchmark_tail)
    leading_benchmarks = [
        benchmark for benchmark in benchmark_order if benchmark not in benchmark_tail_set
    ]
    trailing_benchmarks = [
        benchmark for benchmark in benchmark_tail if benchmark in benchmark_order
    ]
    return leading_benchmarks + trailing_benchmarks


def get_strategy_sort_key(strategy_name: str) -> tuple[int, str]:
    if strategy_name.startswith("kv_score"):
        return STRATEGY_DISPLAY_ORDER["kv_score"], strategy_name
    return STRATEGY_DISPLAY_ORDER.get(strategy_name, len(STRATEGY_DISPLAY_ORDER)), strategy_name


def format_method_label(
    row_label: str,
    token_aliases: dict[str, str],
    variant_order_lookup: dict[frozenset[str], tuple[int, tuple[str, ...]]],
) -> str:
    core_label, annotation = split_method_annotation(row_label)
    strategy_name, ratio_value, variant_suffix = parse_row_label(core_label)
    if ratio_value is None:
        return row_label

    aliased_tokens = apply_method_token_aliases(
        split_variant_tokens(variant_suffix),
        token_aliases,
    )
    variant_order_entry = resolve_variant_order_entry(aliased_tokens, variant_order_lookup)
    display_tokens = list(variant_order_entry[1]) if variant_order_entry else aliased_tokens

    display_label = f"{strategy_name}:{format_ratio_percent(ratio_value)}"
    if display_tokens:
        display_label = f"{display_label}-{'-'.join(display_tokens)}"
    return f"{display_label}{annotation}"


def build_method_sort_key(
    row_label: str,
    token_aliases: dict[str, str],
    variant_order: tuple[tuple[str, ...], ...],
    variant_order_lookup: dict[frozenset[str], tuple[int, tuple[str, ...]]],
    variant_token_priority: dict[str, int],
) -> tuple[object, ...]:
    core_label, annotation = split_method_annotation(row_label)
    special_rank = SPECIAL_METHOD_ORDER.get(core_label)
    if special_rank is not None:
        return (special_rank, 0.0, 0, "", 0, (), "", annotation)

    strategy_name, ratio_value, variant_suffix = parse_row_label(core_label)
    if ratio_value is None:
        return (
            len(SPECIAL_METHOD_ORDER) + 1,
            extract_ratio_from_label(core_label),
            len(STRATEGY_DISPLAY_ORDER),
            core_label,
            0,
            (),
            annotation,
        )

    aliased_tokens = apply_method_token_aliases(
        split_variant_tokens(variant_suffix),
        token_aliases,
    )
    variant_order_entry = resolve_variant_order_entry(aliased_tokens, variant_order_lookup)
    if variant_order_entry is None:
        variant_index = len(variant_order)
        fallback_tokens = tuple(
            sorted(
                aliased_tokens,
                key=lambda token: (
                    variant_token_priority.get(token, len(variant_token_priority)),
                    token,
                ),
            )
        )
    else:
        variant_index, canonical_tokens = variant_order_entry
        fallback_tokens = canonical_tokens

    strategy_rank, strategy_sort_name = get_strategy_sort_key(strategy_name)
    return (
        len(SPECIAL_METHOD_ORDER),
        ratio_value,
        strategy_rank,
        strategy_sort_name,
        variant_index,
        fallback_tokens,
        annotation,
    )


def resolve_case_suffix(case_name: str) -> str:
    return CASE_SUFFIX_ALIASES.get(case_name, case_name)


def build_summary_filename(tag: str) -> str:
    return f"sweep_summary_{tag}.json"


def resolve_model_summary_base_dir(benchmark_root: Path, model_output_dir: str) -> Path:
    return benchmark_root.resolve() / "output" / model_output_dir / "_summaries"


def resolve_shared_summary_base_dir(args: argparse.Namespace) -> Path:
    if args.shared_summary_root is not None:
        return args.shared_summary_root.resolve()
    return args.benchmark_root.resolve() / "output" / "_summaries"


def resolve_tag_summary_path(
    summary_base_dir: Path,
    tag: str,
    candidate_dirs: list[tuple[str, ...]],
) -> Path:
    summary_name = build_summary_filename(tag)
    attempted_paths: list[Path] = []
    for candidate_dir in candidate_dirs:
        path = summary_base_dir / tag
        for component in candidate_dir:
            path /= component
        path /= summary_name
        attempted_paths.append(path)
        if path.is_file():
            return path

    attempted_text = "\n".join(f"- {path}" for path in attempted_paths)
    raise FileNotFoundError(
        f"Could not locate summary JSON for tag {tag!r}. Tried:\n{attempted_text}"
    )


def try_resolve_tag_summary_path(
    summary_base_dir: Path,
    tag: str,
    candidate_dirs: list[tuple[str, ...]],
) -> Path | None:
    if not tag:
        return None
    try:
        return resolve_tag_summary_path(summary_base_dir, tag, candidate_dirs)
    except FileNotFoundError:
        return None


def resolve_baseline_full_summary_path(
    summary_base_dir: Path,
    baseline_tag: str,
    image_priori_mode: str,
) -> Path | None:
    if not baseline_tag:
        return None
    return resolve_tag_summary_path(
        summary_base_dir=summary_base_dir,
        tag=baseline_tag,
        candidate_dirs=[("full",), (image_priori_mode, "full")],
    )


def resolve_baseline_segment_summary_path(
    summary_base_dir: Path,
    baseline_tag: str,
    image_priori_mode: str,
    prefill_mode: str,
) -> Path | None:
    if not baseline_tag:
        return None
    return resolve_tag_summary_path(
        summary_base_dir=summary_base_dir,
        tag=baseline_tag,
        candidate_dirs=[(prefill_mode,), (image_priori_mode, prefill_mode)],
    )


def resolve_optional_method_summary_path(
    summary_base_dir: Path,
    tag: str,
    image_priori_mode: str,
    prefill_mode: str,
) -> Path | None:
    return try_resolve_tag_summary_path(
        summary_base_dir=summary_base_dir,
        tag=tag,
        candidate_dirs=[(image_priori_mode, prefill_mode), (prefill_mode,)],
    )


def resolve_shared_none_summary_path(
    shared_summary_base_dir: Path,
    tag: str,
    none_image_priori_mode: str,
    prefill_mode: str,
) -> Path | None:
    return try_resolve_tag_summary_path(
        summary_base_dir=shared_summary_base_dir,
        tag=tag,
        candidate_dirs=[
            (none_image_priori_mode, prefill_mode),
            (prefill_mode,),
            (none_image_priori_mode,),
            (),
        ],
    )


def resolve_case_summary_paths(
    args: argparse.Namespace,
    summary_base_dir: Path,
) -> list[Path]:
    if args.case_summary:
        return [path.resolve() for path in args.case_summary]

    paths: list[Path] = []
    for case_name in split_csv(args.cases):
        case_suffix = resolve_case_suffix(case_name)
        if case_suffix == "":
            case_tag = args.experiment_tag
        else:
            case_tag = f"{args.experiment_tag}_{case_suffix}"
        path = resolve_tag_summary_path(
            summary_base_dir=summary_base_dir,
            tag=case_tag,
            candidate_dirs=[
                (args.image_priori_mode, args.prefill_mode),
                (args.prefill_mode,),
            ],
        )
        paths.append(path)
    return paths


def model_output_sort_key(model_output_dir: str) -> tuple[int, int | str]:
    try:
        return 0, PREFERRED_MODEL_ORDER.index(model_output_dir)
    except ValueError:
        return 1, model_output_dir


def unique_preserve_order(values: list[str]) -> list[str]:
    unique_values: list[str] = []
    seen: set[str] = set()
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        unique_values.append(value)
    return unique_values


def resolve_case_tags(args: argparse.Namespace) -> list[str]:
    if args.case_summary:
        return []

    case_tags: list[str] = []
    for case_name in split_csv(args.cases):
        case_suffix = resolve_case_suffix(case_name)
        if case_suffix == "":
            case_tag = args.experiment_tag
        else:
            case_tag = f"{args.experiment_tag}_{case_suffix}"
        if case_tag:
            case_tags.append(case_tag)
    return unique_preserve_order(case_tags)


def collect_model_discovery_tags(
    args: argparse.Namespace,
    full_baseline_tag: str,
    segment_baseline_tag: str,
) -> list[str]:
    return unique_preserve_order(
        [
            tag
            for tag in [
                full_baseline_tag,
                segment_baseline_tag,
                *resolve_case_tags(args),
                args.infoblend_tag.strip(),
                args.cacheblend_tag.strip(),
                args.kvshare_tag.strip(),
            ]
            if tag
        ]
    )


def discover_model_output_dirs(
    benchmark_root: Path,
    summary_tags: list[str],
) -> list[str]:
    output_root = benchmark_root.resolve() / "output"
    if not output_root.is_dir():
        raise NotADirectoryError(f"Output directory does not exist: {output_root}")

    if not summary_tags:
        raise ValueError(
            "Could not auto-discover model output directories because no model-scoped "
            "summary tags are configured. Set at least one non-empty baseline/method "
            "tag, or pass --model-output-dir explicitly."
        )

    discovered: list[str] = []
    for child in sorted(output_root.iterdir()):
        if not child.is_dir() or child.name == "_summaries":
            continue
        summary_base_dir = child / "_summaries"
        if not summary_base_dir.is_dir():
            continue
        if not any((summary_base_dir / tag).exists() for tag in summary_tags):
            continue
        discovered.append(child.name)

    if not discovered:
        raise FileNotFoundError(
            "Could not discover any model output directories under "
            f"{output_root} that contain any of these tags: {summary_tags!r}."
        )

    return sorted(discovered, key=model_output_sort_key)


def resolve_model_output_dirs(
    args: argparse.Namespace,
    summary_tags: list[str],
) -> list[str]:
    requested = unique_preserve_order(
        [item for value in args.model_output_dir for item in split_csv(value)]
    )
    if not requested:
        return discover_model_output_dirs(
            benchmark_root=args.benchmark_root,
            summary_tags=summary_tags,
        )

    for model_output_dir in requested:
        summary_base_dir = resolve_model_summary_base_dir(args.benchmark_root, model_output_dir)
        if not summary_base_dir.is_dir():
            raise FileNotFoundError(
                f"Model summary directory does not exist for {model_output_dir!r}: "
                f"{summary_base_dir}"
            )
    return requested


def resolve_output_path(
    args: argparse.Namespace,
    model_output_dirs: list[str],
) -> Path:
    if args.output is not None:
        return args.output.resolve()

    if args.model_output_dir and len(model_output_dirs) == 1:
        return (
            resolve_model_summary_base_dir(args.benchmark_root, model_output_dirs[0])
            / args.experiment_tag
            / "benchmark_metric_table.csv"
        ).resolve()

    return (
        args.benchmark_root.resolve()
        / "output"
        / "_summaries"
        / args.experiment_tag
        / "benchmark_metric_table.csv"
    ).resolve()


def resolve_xlsx_output_path(args: argparse.Namespace, csv_output_path: Path) -> Path:
    if args.xlsx_output is not None:
        return args.xlsx_output.resolve()
    return csv_output_path.with_suffix(".xlsx")


def canonicalize_model_part(value: str) -> str:
    return NON_ALNUM_PATTERN.sub("", value.lower())


def value_matches_model_output_dir(value: object, model_output_dir: str) -> bool:
    if value in (None, ""):
        return False

    target = canonicalize_model_part(model_output_dir)
    normalized = str(value).replace("\\", "/").strip("/")
    parts = [part for part in normalized.split("/") if part]
    return any(canonicalize_model_part(part) == target for part in parts)


def record_matches_model_output_dir(record: dict, model_output_dir: str) -> bool:
    for key in ("model_label", "model_name_or_path", "output_dir"):
        if value_matches_model_output_dir(record.get(key), model_output_dir):
            return True
    return False


def summary_is_multi_model(summary: dict) -> bool:
    sweep_config = summary.get("sweep_config") or {}
    model_name_or_paths = [
        value
        for value in (sweep_config.get("model_name_or_paths") or [])
        if value not in (None, "")
    ]
    return len(model_name_or_paths) > 1


def summary_matches_model_output_dir(summary: dict, model_output_dir: str) -> bool:
    sweep_config = summary.get("sweep_config") or {}
    candidate_values = [
        summary.get("summary_json"),
        sweep_config.get("model_name_or_path"),
        *(sweep_config.get("model_name_or_paths") or []),
    ]
    return any(
        value_matches_model_output_dir(candidate_value, model_output_dir)
        for candidate_value in candidate_values
    )


def get_summary_records_for_model(
    summary: dict,
    model_output_dir: str | None = None,
) -> list[dict]:
    runs = [record for record in (summary.get("runs") or []) if should_include_record(record)]
    summary_rows = [
        record
        for record in (summary.get("summary_rows") or [])
        if should_include_record(record)
    ]

    if model_output_dir is None:
        return runs or summary_rows

    filtered_runs = [
        record for record in runs if record_matches_model_output_dir(record, model_output_dir)
    ]
    if filtered_runs:
        return filtered_runs

    filtered_summary_rows = [
        record
        for record in summary_rows
        if record_matches_model_output_dir(record, model_output_dir)
    ]
    if filtered_summary_rows:
        return filtered_summary_rows

    if summary_is_multi_model(summary):
        return []

    if summary_matches_model_output_dir(summary, model_output_dir):
        return runs or summary_rows

    return []


def remap_row_label_strategy_name(
    row_label: str,
    strategy_name_aliases: dict[str, str],
) -> str:
    if not strategy_name_aliases:
        return row_label

    core_label, annotation = split_method_annotation(row_label)
    strategy_name, ratio_value, variant_suffix = parse_row_label(core_label)
    if ratio_value is None:
        return row_label

    mapped_name = strategy_name_aliases.get(strategy_name)
    if not mapped_name or mapped_name == strategy_name:
        return row_label

    mapped_core = f"{mapped_name}:{format_ratio_percent(ratio_value)}"
    if variant_suffix:
        mapped_core = f"{mapped_core}-{variant_suffix}"
    return f"{mapped_core}{annotation}"


def is_plain_vanilla_record(record: dict) -> bool:
    recompute_strategy = str(record.get("recompute_strategy", "none"))
    return recompute_strategy == "none"


def collect_rows_from_summary(
    summary: dict,
    source_name: str,
    strategy_name_aliases: dict[str, str] | None = None,
    record_filter: Callable[[dict], bool] | None = None,
    model_output_dir: str | None = None,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for record in get_summary_records_for_model(summary, model_output_dir=model_output_dir):
        if record_filter is not None and not record_filter(record):
            continue

        recompute_strategy = str(record.get("recompute_strategy", "none"))
        ratio_percent = (
            float(record["ratio_percent"])
            if record.get("ratio_percent") is not None
            else None
        )
        row_label = normalize_row_label(
            source_name=source_name,
            recompute_strategy=recompute_strategy,
            ratio_percent=ratio_percent,
            run_name=record.get("run_name"),
        )
        if strategy_name_aliases:
            row_label = remap_row_label_strategy_name(row_label, strategy_name_aliases)

        benchmark = get_benchmark_name(record)
        metric_name, metric_value = get_metric_entry(record)
        rows.append(
            {
                "source_name": source_name,
                "row_label": row_label,
                "benchmark": benchmark,
                "metric_name": metric_name,
                "metric_value": metric_value,
                "image_priori_mode": get_record_image_priori_mode(summary, record),
                "config_entry_label": record.get("config_entry_label"),
                "test_file": get_record_test_file(record),
                "input_max_length": get_record_input_max_length(record),
                "run_name": record.get("run_name"),
                "score_file": record.get("score_file"),
            }
        )
    return rows


def collect_segment_baseline_rows(
    summary: dict,
    model_output_dir: str,
) -> list[dict[str, object]]:
    return collect_rows_from_summary(
        summary=summary,
        source_name="image_segment",
        record_filter=is_plain_vanilla_record,
        model_output_dir=model_output_dir,
    )


def collect_shared_none_rows(summary: dict, model_output_dir: str) -> list[dict[str, object]]:
    return collect_rows_from_summary(
        summary=summary,
        source_name="none",
        record_filter=lambda record: is_plain_vanilla_record(record)
        and record_matches_model_output_dir(record, model_output_dir),
        model_output_dir=model_output_dir,
    )


def get_model_display_name(
    model_output_dir: str,
    candidate_summaries: list[dict | None],
) -> str:
    for summary in candidate_summaries:
        if summary is None:
            continue

        sweep_config = summary.get("sweep_config") or {}
        model_name_or_path = sweep_config.get("model_name_or_path")
        if model_name_or_path not in (None, ""):
            if not summary_is_multi_model(summary) or value_matches_model_output_dir(
                model_name_or_path,
                model_output_dir,
            ):
                return Path(str(model_name_or_path)).name

        for candidate_path in sweep_config.get("model_name_or_paths") or []:
            if value_matches_model_output_dir(candidate_path, model_output_dir):
                return Path(str(candidate_path)).name

    return model_output_dir


def collect_global_benchmark_metadata(
    records_by_model: list[list[dict[str, object]]],
    benchmark_tail: tuple[str, ...],
) -> tuple[list[str], dict[str, str]]:
    benchmark_order: list[str] = []
    benchmark_metrics: dict[str, str] = {}

    for records in records_by_model:
        for record in records:
            benchmark = str(record["benchmark"])
            metric_name = str(record["metric_name"])
            if benchmark not in benchmark_order:
                benchmark_order.append(benchmark)

            existing_metric = benchmark_metrics.get(benchmark)
            if existing_metric is None:
                benchmark_metrics[benchmark] = metric_name
            elif existing_metric != metric_name:
                raise ValueError(
                    f"Benchmark {benchmark!r} has inconsistent metric names: "
                    f"{existing_metric!r} vs {metric_name!r}"
                )

    return reorder_benchmarks(benchmark_order, benchmark_tail), benchmark_metrics


def build_table_from_records(
    records: list[dict[str, object]],
    token_aliases: dict[str, str],
    variant_order: tuple[tuple[str, ...], ...],
    benchmark_tail: tuple[str, ...],
    ordered_benchmarks: list[str] | None = None,
    benchmark_metrics_override: dict[str, str] | None = None,
) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    benchmark_order: list[str] = []
    benchmark_metrics: dict[str, str] = {}
    table: dict[str, dict[str, float]] = {}
    variant_order_lookup = build_variant_order_lookup(variant_order)
    variant_token_priority = build_variant_token_priority(variant_order)

    for record in records:
        row_label = str(record["row_label"])
        benchmark = str(record["benchmark"])
        metric_name = str(record["metric_name"])
        metric_value_raw = record["metric_value"]
        if not isinstance(metric_value_raw, (int, float, str)):
            raise TypeError(
                "metric_value must be int, float, or str, got "
                f"{type(metric_value_raw).__name__} for row {row_label!r}"
            )
        metric_value = float(cast(int | float | str, metric_value_raw))

        if benchmark not in benchmark_order:
            benchmark_order.append(benchmark)

        existing_metric = benchmark_metrics.get(benchmark)
        if existing_metric is None:
            benchmark_metrics[benchmark] = metric_name
        elif existing_metric != metric_name:
            raise ValueError(
                f"Benchmark {benchmark!r} has inconsistent metric names: "
                f"{existing_metric!r} vs {metric_name!r}"
            )

        row_values = table.setdefault(row_label, {})
        if benchmark in row_values:
            raise ValueError(
                f"Duplicate value for row={row_label!r}, benchmark={benchmark!r}"
            )
        row_values[benchmark] = metric_value

    if ordered_benchmarks is None:
        ordered_benchmarks = reorder_benchmarks(benchmark_order, benchmark_tail)
    else:
        missing_benchmarks = [
            benchmark for benchmark in benchmark_order if benchmark not in ordered_benchmarks
        ]
        if missing_benchmarks:
            raise ValueError(
                "The provided ordered_benchmarks list is missing benchmarks present "
                f"in the records: {missing_benchmarks}"
            )

    effective_benchmark_metrics = dict(benchmark_metrics)
    if benchmark_metrics_override is not None:
        for benchmark, metric_name in benchmark_metrics.items():
            override_metric_name = benchmark_metrics_override.get(benchmark)
            if override_metric_name is None:
                raise ValueError(
                    f"Missing metric name for benchmark {benchmark!r} in benchmark_metrics_override"
                )
            if override_metric_name != metric_name:
                raise ValueError(
                    f"Benchmark {benchmark!r} has inconsistent metric names: "
                    f"{metric_name!r} vs {override_metric_name!r}"
                )
        effective_benchmark_metrics = dict(benchmark_metrics_override)

    headers = [
        "method",
        *[
            f"{benchmark} ({effective_benchmark_metrics[benchmark]})"
            for benchmark in ordered_benchmarks
        ],
    ]
    rows: list[dict[str, str]] = []
    display_labels: dict[str, str] = {}
    for row_label in sorted(
        table,
        key=lambda label: build_method_sort_key(
            label,
            token_aliases=token_aliases,
            variant_order=variant_order,
            variant_order_lookup=variant_order_lookup,
            variant_token_priority=variant_token_priority,
        ),
    ):
        display_label = format_method_label(
            row_label,
            token_aliases=token_aliases,
            variant_order_lookup=variant_order_lookup,
        )
        existing_label = display_labels.get(display_label)
        if existing_label is not None and existing_label != row_label:
            raise ValueError(
                "Method label aliases produced duplicate output row labels: "
                f"{existing_label!r} and {row_label!r} both map to {display_label!r}"
            )
        display_labels[display_label] = row_label

        row = {"method": display_label}
        for benchmark in ordered_benchmarks:
            header = f"{benchmark} ({effective_benchmark_metrics[benchmark]})"
            value = table[row_label].get(benchmark)
            row[header] = "" if value is None else format_value(value)
        rows.append(row)

    return headers, rows, effective_benchmark_metrics


def build_combined_sheet_rows(
    model_sections: list[tuple[str, list[dict[str, str]]]],
    headers: list[str],
) -> list[list[str]]:
    row_length = len(headers)
    blank_row = [""] * row_length
    output_rows: list[list[str]] = []

    for index, (model_name, section_rows) in enumerate(model_sections):
        if index > 0:
            output_rows.append(blank_row.copy())

        output_rows.append([f"model: {model_name}", *([""] * (row_length - 1))])
        output_rows.append(list(headers))
        for row in section_rows:
            output_rows.append([row.get(header, "") for header in headers])

    return output_rows


def write_block_csv(output_path: Path, rows: list[list[str]]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)


def write_xlsx(output_path: Path, rows: list[list[str]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to write the XLSX export. Install it with pip install openpyxl."
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "benchmark_metric_table"

    model_font = Font(bold=True, size=12)
    model_fill = PatternFill(fill_type="solid", fgColor="EDEDED")
    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
    centered = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=1):
        worksheet.append(row)
        if row and str(row[0]).startswith("model: "):
            cell = worksheet.cell(row=row_index, column=1)
            cell.font = model_font
            cell.fill = model_fill
        elif row and row[0] == "method":
            for column_index, value in enumerate(row, start=1):
                if value == "":
                    continue
                cell = worksheet.cell(row=row_index, column=column_index)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = centered

    for column_cells in worksheet.columns:
        column_index = column_cells[0].column
        max_length = 0
        for cell in column_cells:
            cell_text = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(cell_text))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 10),
            40,
        )

    workbook.save(output_path)


def main() -> None:
    args = parse_args()
    token_aliases = parse_assignment_csv(args.method_token_aliases)
    benchmark_tail = tuple(split_csv(args.benchmark_tail))
    variant_order = parse_variant_order_csv(args.method_variant_order)
    full_baseline_tag = resolve_configured_summary_tag(
        args.full_baseline_tag,
        args.baseline_tag,
    )
    segment_baseline_tag = resolve_configured_summary_tag(
        args.segment_baseline_tag,
        args.baseline_tag,
    )
    model_discovery_tags = collect_model_discovery_tags(
        args,
        full_baseline_tag=full_baseline_tag,
        segment_baseline_tag=segment_baseline_tag,
    )
    model_output_dirs = resolve_model_output_dirs(args, model_discovery_tags)
    if args.case_summary and len(model_output_dirs) != 1:
        raise ValueError(
            "Explicit --case-summary paths are only supported when exporting exactly one model."
        )

    shared_summary_base_dir = resolve_shared_summary_base_dir(args)
    shared_none_summary_path = resolve_shared_none_summary_path(
        shared_summary_base_dir=shared_summary_base_dir,
        tag=args.none_tag,
        none_image_priori_mode=args.none_image_priori_mode,
        prefill_mode=args.prefill_mode,
    )
    shared_none_summary = (
        load_json(shared_none_summary_path) if shared_none_summary_path is not None else None
    )

    model_blocks: list[dict[str, object]] = []
    records_by_model: list[list[dict[str, object]]] = []

    for model_output_dir in model_output_dirs:
        summary_base_dir = resolve_model_summary_base_dir(args.benchmark_root, model_output_dir)
        full_summary_path = resolve_baseline_full_summary_path(
            summary_base_dir=summary_base_dir,
            baseline_tag=full_baseline_tag,
            image_priori_mode=args.image_priori_mode,
        )
        segment_summary_path = resolve_baseline_segment_summary_path(
            summary_base_dir=summary_base_dir,
            baseline_tag=segment_baseline_tag,
            image_priori_mode=args.image_priori_mode,
            prefill_mode=args.prefill_mode,
        )
        case_summary_paths = resolve_case_summary_paths(args, summary_base_dir)
        infoblend_summary_path = resolve_optional_method_summary_path(
            summary_base_dir=summary_base_dir,
            tag=args.infoblend_tag,
            image_priori_mode=args.image_priori_mode,
            prefill_mode=args.prefill_mode,
        )
        cacheblend_summary_path = resolve_optional_method_summary_path(
            summary_base_dir=summary_base_dir,
            tag=args.cacheblend_tag,
            image_priori_mode=args.image_priori_mode,
            prefill_mode=args.prefill_mode,
        )
        kvshare_summary_path = resolve_optional_method_summary_path(
            summary_base_dir=summary_base_dir,
            tag=args.kvshare_tag,
            image_priori_mode=args.image_priori_mode,
            prefill_mode=args.prefill_mode,
        )

        full_summary = load_json(full_summary_path) if full_summary_path is not None else None
        segment_summary = (
            load_json(segment_summary_path) if segment_summary_path is not None else None
        )
        case_summaries = [load_json(path) for path in case_summary_paths]
        infoblend_summary = (
            load_json(infoblend_summary_path) if infoblend_summary_path is not None else None
        )
        cacheblend_summary = (
            load_json(cacheblend_summary_path) if cacheblend_summary_path is not None else None
        )
        kvshare_summary = (
            load_json(kvshare_summary_path) if kvshare_summary_path is not None else None
        )

        combined_records = consolidate_export_records(
            (
                collect_rows_from_summary(
                    full_summary,
                    "full",
                    model_output_dir=model_output_dir,
                )
                if full_summary is not None
                else []
            )
            + (
                collect_segment_baseline_rows(segment_summary, model_output_dir)
                if segment_summary is not None
                else []
            )
            + (
                collect_shared_none_rows(shared_none_summary, model_output_dir)
                if shared_none_summary is not None
                else []
            )
            + (
                collect_rows_from_summary(
                    infoblend_summary,
                    "infoblend",
                    strategy_name_aliases={"first": "infoblend"},
                    model_output_dir=model_output_dir,
                )
                if infoblend_summary is not None
                else []
            )
            + (
                collect_rows_from_summary(
                    cacheblend_summary,
                    "cacheblend",
                    model_output_dir=model_output_dir,
                )
                if cacheblend_summary is not None
                else []
            )
            + (
                collect_rows_from_summary(
                    kvshare_summary,
                    "kvshare",
                    model_output_dir=model_output_dir,
                )
                if kvshare_summary is not None
                else []
            )
            + [
                record
                for summary in case_summaries
                for record in collect_rows_from_summary(
                    summary,
                    "image_segment",
                    model_output_dir=model_output_dir,
                )
            ]
        )

        model_display_name = get_model_display_name(
            model_output_dir,
            [
                full_summary,
                segment_summary,
                *case_summaries,
                infoblend_summary,
                cacheblend_summary,
                kvshare_summary,
                shared_none_summary,
            ],
        )
        model_blocks.append(
            {
                "model_output_dir": model_output_dir,
                "model_display_name": model_display_name,
                "full_summary_path": full_summary_path,
                "segment_summary_path": segment_summary_path,
                "case_summary_paths": case_summary_paths,
                "infoblend_summary_path": infoblend_summary_path,
                "cacheblend_summary_path": cacheblend_summary_path,
                "kvshare_summary_path": kvshare_summary_path,
                "records": combined_records,
            }
        )
        records_by_model.append(combined_records)

    ordered_benchmarks, benchmark_metrics = collect_global_benchmark_metadata(
        records_by_model=records_by_model,
        benchmark_tail=benchmark_tail,
    )

    headers = [
        "method",
        *[
            f"{benchmark} ({benchmark_metrics[benchmark]})"
            for benchmark in ordered_benchmarks
        ],
    ]
    model_sections: list[tuple[str, list[dict[str, str]]]] = []
    for block in model_blocks:
        _, rows, _ = build_table_from_records(
            cast(list[dict[str, object]], block["records"]),
            token_aliases=token_aliases,
            variant_order=variant_order,
            benchmark_tail=benchmark_tail,
            ordered_benchmarks=ordered_benchmarks,
            benchmark_metrics_override=benchmark_metrics,
        )
        model_sections.append((str(block["model_display_name"]), rows))

    csv_output_path = resolve_output_path(args, model_output_dirs)
    xlsx_output_path = resolve_xlsx_output_path(args, csv_output_path)
    combined_sheet_rows = build_combined_sheet_rows(model_sections, headers)

    write_block_csv(csv_output_path, combined_sheet_rows)
    write_xlsx(xlsx_output_path, combined_sheet_rows)

    print(f"Wrote CSV to: {csv_output_path}")
    print(f"Wrote XLSX to: {xlsx_output_path}")
    if shared_none_summary_path is not None:
        print(f"Shared none summary: {shared_none_summary_path}")
    print("Models:")
    for block in model_blocks:
        print(f"  {block['model_display_name']} ({block['model_output_dir']})")
        full_summary_path = block["full_summary_path"]
        if full_summary_path is not None:
            print(f"    Full baseline summary: {full_summary_path}")
        segment_summary_path = block["segment_summary_path"]
        if segment_summary_path is not None:
            print(f"    Image-segment baseline summary: {segment_summary_path}")
        infoblend_summary_path = block["infoblend_summary_path"]
        if infoblend_summary_path is not None:
            print(f"    Infoblend summary: {infoblend_summary_path}")
        cacheblend_summary_path = block["cacheblend_summary_path"]
        if cacheblend_summary_path is not None:
            print(f"    Cacheblend summary: {cacheblend_summary_path}")
        kvshare_summary_path = block["kvshare_summary_path"]
        if kvshare_summary_path is not None:
            print(f"    KVShare summary: {kvshare_summary_path}")
        for path in cast(list[Path], block["case_summary_paths"]):
            print(f"    KV score case summary: {path}")
    print("Benchmark metrics:")
    for benchmark in ordered_benchmarks:
        print(f"  {benchmark}: {benchmark_metrics[benchmark]}")


if __name__ == "__main__":
    main()
